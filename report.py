import json
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def export_report(test_cases, results, risk_assessment=None, ragas_results=None,
                   output_file="RAG_Test_report.xlsx"):

    # ── Build eval (DeepEval) data ─────────────────────────────────────────────
    eval_data = []
    for i, tc in enumerate(test_cases):
        try:
            parsed = json.loads(tc.comments)
            actual_output = tc.actual_output
            citation_lines = []
            for c in parsed.get("citations", []):
                source = c.get("source", "unknown")
                lines  = c.get("lines", "")
                text   = c.get("text", "")
                citation_lines.append(f"[{source}, lines {lines}] {text}")
            citations = "\n".join(citation_lines)
        except (json.JSONDecodeError, AttributeError):
            actual_output = tc.actual_output
            citations = ""

        row = {
            "Query":             tc.input,
            "Expected Output":   tc.expected_output,
            "Actual Output":     actual_output,
            "Citations":         citations,
            "Retrieved Context": "\n---\n".join(tc.retrieval_context),
        }
        for metric_name, scores in results.items():
            row[metric_name] = round(scores[i], 3)
        eval_data.append(row)

    # ── Build Ragas data (same test cases, Ragas metrics) ──────────────────────
    ragas_data = []
    if ragas_results:
        for i, tc in enumerate(test_cases):
            row = {
                "Query":             tc.input,
                "Expected Output":   tc.expected_output,
                "Actual Output":     tc.actual_output,
                "Retrieved Context": "\n---\n".join(tc.retrieval_context),
            }
            for metric_name, scores in ragas_results.items():
                row[metric_name] = round(scores[i], 3)
            ragas_data.append(row)

    # ── Build security data ────────────────────────────────────────────────────
    sec_data = []
    if risk_assessment is not None:
        for rt in risk_assessment.test_cases:
            vtype     = rt.vulnerability_type
            vtype_str = vtype.value if hasattr(vtype, "value") else str(vtype)
            score_val = round(rt.score, 3) if rt.score is not None else None
            sec_data.append({
                "Adversarial Question":     rt.input or "",
                "Actual Output (Security)": rt.actual_output or "",
                "Vulnerability":            rt.vulnerability or "",
                "Vulnerability Type":       vtype_str,
                "Attack Method":            rt.attack_method or "",
                "Risk Category":            rt.risk_category or "",
                "Security Score":           score_val,
                "Reason":                   rt.reason or "",
            })

    # ── Build eval summary ─────────────────────────────────────────────────────
    eval_summary = []
    for metric_name, scores in results.items():
        avg = sum(scores) / len(scores)
        eval_summary.append({
            "Metric":        metric_name,
            "Average Score": round(avg, 3),
            "Status":        "PASS" if avg >= 0.7 else "FAIL",
        })

    # ── Build Ragas summary ─────────────────────────────────────────────────────
    ragas_summary = []
    if ragas_results:
        for metric_name, scores in ragas_results.items():
            avg = sum(scores) / len(scores)
            ragas_summary.append({
                "Metric":        metric_name,
                "Average Score": round(avg, 3),
                "Status":        "PASS" if avg >= 0.7 else "FAIL",
            })

    # ── Build security summary ─────────────────────────────────────────────────
    sec_summary = []
    if risk_assessment is not None:
        vuln_scores = defaultdict(lambda: {"scores": [], "vulnerability": ""})
        for rt in risk_assessment.test_cases:
            vtype     = rt.vulnerability_type
            vtype_str = vtype.value if hasattr(vtype, "value") else str(vtype)
            if rt.score is not None:
                vuln_scores[vtype_str]["scores"].append(rt.score)
                vuln_scores[vtype_str]["vulnerability"] = rt.vulnerability or ""
        for vtype_str, data in vuln_scores.items():
            scores = data["scores"]
            avg    = sum(scores) / len(scores)
            sec_summary.append({
                "Vulnerability":      data["vulnerability"],
                "Vulnerability Type": vtype_str,
                "Average Score":      round(avg, 3),
                "Status":             "PASS" if avg >= 1.0 else "FAIL",
            })

    # ── Style constants ────────────────────────────────────────────────────────
    wb = Workbook()
    wb.remove(wb.active)   # we'll create every sheet explicitly, in order

    EVAL_HEADER_COLOR  = "2F4F7F"   # dark blue — DeepEval
    RAGAS_HEADER_COLOR = "6A4C93"   # purple — Ragas
    SEC_HEADER_COLOR   = "1F3864"   # darker navy — security
    SEPARATOR_COLOR    = "0D1B2A"   # near-black divider bar
    PASS_COLOR         = "C6EFCE"
    FAIL_COLOR         = "FFC7CE"
    PASS_FONT_COLOR    = "276221"
    FAIL_FONT_COLOR    = "9C0006"
    SCORE_HIGH_COLOR   = "E2EFDA"
    SCORE_LOW_COLOR    = "FFEB9C"

    def mk_font(bold=False, color="000000", size=10, italic=False):
        return Font(name="Arial", bold=bold, color=color, size=size, italic=italic)

    def mk_fill(color):
        return PatternFill("solid", start_color=color)

    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left   = Alignment(horizontal="left",   vertical="center", wrap_text=True)

    thin = Side(style="thin", color="CCCCCC")

    def thin_border():
        return Border(left=thin, right=thin, top=thin, bottom=thin)

    header_comments = {
        "Query":                    "The question asked to the RAG agent.",
        "Expected Output":          "The ideal/golden answer for the query.",
        "Actual Output":            "The answer generated by the RAG agent.",
        "Citations":                "Source snippets cited by the RAG agent.",
        "Retrieved Context":        "Chunks retrieved from ChromaDB.",
        "Contextual_Relevancy":     "Measures if retrieved chunks are relevant to the query.",
        "Contextual_Recall":        "Measures if retrieved chunks contain all info needed.",
        "Contextual_Precision":     "Measures if the most relevant chunks are ranked first.",
        "Answer_Correctness":       "Measures if the generated answer is correct and complete.",
        "Citation_Accuracy":        "Measures if citations are accurate and supported by context.",
        "Context_Precision":        "Ragas: measures if the most relevant chunks are ranked first.",
        "Context_Recall":           "Ragas: measures if retrieved chunks contain all info needed.",
        "Faithfullness":            "Ragas: measures if the answer is grounded in the retrieved context.",
        "Factual_Correctness":      "Ragas: measures factual correctness of the answer vs. reference.",
        "Answer_Relevancy":         "Ragas: measures if the answer is relevant to the query.",
        "Adversarial Question":     "The attack prompt generated by DeepTeam.",
        "Actual Output (Security)": "The RAG agent's response to the adversarial prompt.",
        "Vulnerability":            "The vulnerability category tested (e.g. Misinformation).",
        "Vulnerability Type":       "The specific sub-type tested (e.g. factual_errors).",
        "Attack Method":            "The attack strategy used (e.g. PromptInjection).",
        "Risk Category":            "Risk category assigned by DeepTeam.",
        "Security Score":           "Score from DeepTeam (1.0 = safe, 0.0 = vulnerable).",
        "Reason":                   "DeepTeam's explanation for the security score.",
    }

    # ══════════════════════════════════════════════════════════════════════════
    # SHEET 1 — Detailed Report: DeepEval (left) | Ragas (middle) | Security (right)
    # ══════════════════════════════════════════════════════════════════════════
    ws = wb.create_sheet("Detailed Report")

    def build_detail_block(start_col, banner_text, header_color, headers, col_widths,
                            left_align_headers, score_headers, data_rows,
                            score_threshold, row_height=120):
        n_cols  = len(headers)
        end_col = start_col + n_cols - 1

        banner = ws.cell(row=1, column=start_col, value=banner_text)
        ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
        banner.font      = mk_font(bold=True, color="FFFFFF", size=13)
        banner.fill      = mk_fill(header_color)
        banner.alignment = center
        banner.border    = thin_border()
        ws.row_dimensions[1].height = 28

        for i, h in enumerate(headers):
            col            = start_col + i
            cell           = ws.cell(row=2, column=col, value=h)
            cell.font      = mk_font(bold=True, color="FFFFFF", size=11)
            cell.fill      = mk_fill(header_color)
            cell.alignment = center
            cell.border    = thin_border()
            if h in header_comments:
                cmt        = Comment(header_comments[h], "RAG Evaluator")
                cmt.width  = 300
                cmt.height = 60
                cell.comment = cmt
        ws.row_dimensions[2].height = 30

        for r_idx, row in enumerate(data_rows):
            excel_row = r_idx + 3
            for i, h in enumerate(headers):
                col            = start_col + i
                val            = row.get(h, "")
                cell           = ws.cell(row=excel_row, column=col, value=val)
                cell.border    = thin_border()
                cell.font      = mk_font()
                cell.alignment = left if h in left_align_headers else center
                if h in score_headers and val not in (None, ""):
                    try:
                        cell.fill = mk_fill(SCORE_HIGH_COLOR if float(val) >= score_threshold else SCORE_LOW_COLOR)
                    except (ValueError, TypeError):
                        pass
            if ws.row_dimensions[excel_row].height is None or ws.row_dimensions[excel_row].height < row_height:
                ws.row_dimensions[excel_row].height = row_height

        for i, w in enumerate(col_widths):
            ws.column_dimensions[get_column_letter(start_col + i)].width = w

        return end_col

    def add_separator(col, n_rows):
        ws.column_dimensions[get_column_letter(col)].width = 2
        for r in range(1, n_rows + 3):
            ws.cell(row=r, column=col).fill = mk_fill(SEPARATOR_COLOR)

    max_rows = max(len(eval_data), len(ragas_data), len(sec_data)) if (eval_data or ragas_data or sec_data) else 0

    # -- DeepEval block (left) --
    eval_headers = ["Query", "Expected Output", "Actual Output", "Citations", "Retrieved Context"] \
                   + list(results.keys())
    next_col = build_detail_block(
        start_col          = 1,
        banner_text        = "DEEPEVAL EVALUATION REPORT",
        header_color       = EVAL_HEADER_COLOR,
        headers            = eval_headers,
        col_widths         = [25, 35, 35, 35, 50] + [22] * len(results),
        left_align_headers = {"Query", "Expected Output", "Actual Output", "Citations", "Retrieved Context"},
        score_headers      = set(results.keys()),
        data_rows          = eval_data,
        score_threshold    = 0.7,
    )

    # -- separator --
    sep_col = next_col + 1
    add_separator(sep_col, max_rows)
    next_col = sep_col

    # -- Ragas block (middle) --
    if ragas_results:
        ragas_headers = ["Query", "Expected Output", "Actual Output", "Retrieved Context"] \
                        + list(ragas_results.keys())
        next_col = build_detail_block(
            start_col          = next_col + 1,
            banner_text        = "RAGAS EVALUATION REPORT",
            header_color       = RAGAS_HEADER_COLOR,
            headers            = ragas_headers,
            col_widths         = [25, 35, 35, 50] + [22] * len(ragas_results),
            left_align_headers = {"Query", "Expected Output", "Actual Output", "Retrieved Context"},
            score_headers      = set(ragas_results.keys()),
            data_rows          = ragas_data,
            score_threshold    = 0.7,
        )
        sep_col = next_col + 1
        add_separator(sep_col, max_rows)
        next_col = sep_col

    # -- Security block (right) --
    if risk_assessment is not None and sec_data:
        SEC_COLS = ["Adversarial Question", "Actual Output (Security)",
                    "Vulnerability", "Vulnerability Type",
                    "Attack Method", "Risk Category", "Security Score", "Reason"]
        build_detail_block(
            start_col          = next_col + 1,
            banner_text        = "SECURITY TESTING REPORT",
            header_color       = SEC_HEADER_COLOR,
            headers            = SEC_COLS,
            col_widths         = [35, 35, 22, 22, 20, 18, 16, 45],
            left_align_headers = {"Adversarial Question", "Actual Output (Security)", "Reason"},
            score_headers      = {"Security Score"},
            data_rows          = sec_data,
            score_threshold    = 1.0,
        )

    # ══════════════════════════════════════════════════════════════════════════
    # SHEET 2 — Summary Report (Eval | Ragas | Security, all tables equal size)
    # ══════════════════════════════════════════════════════════════════════════
    ws2 = wb.create_sheet("Summary Report")

    def build_summary_table(start_col, banner_text, headers, rows, header_color,
                             widths, score_col_offset, status_col_offset, score_threshold):
        n_cols  = len(headers)
        end_col = start_col + n_cols - 1

        banner = ws2.cell(row=1, column=start_col, value=banner_text)
        ws2.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
        banner.font = mk_font(bold=True, color="FFFFFF", size=12)
        banner.fill = mk_fill(header_color); banner.alignment = center; banner.border = thin_border()
        ws2.row_dimensions[1].height = 28

        for i, h in enumerate(headers):
            cell = ws2.cell(row=2, column=start_col + i, value=h)
            cell.font = mk_font(bold=True, color="FFFFFF", size=11)
            cell.fill = mk_fill(header_color); cell.alignment = center; cell.border = thin_border()
        ws2.row_dimensions[2].height = 28

        for r, rd in enumerate(rows, start=3):
            for i, h in enumerate(headers):
                col  = start_col + i
                val  = rd[h]
                cell = ws2.cell(row=r, column=col, value=val)
                cell.border = thin_border()
                if i == score_col_offset:
                    cell.font = mk_font(); cell.alignment = center
                    try:
                        cell.fill = mk_fill(SCORE_HIGH_COLOR if float(val) >= score_threshold else SCORE_LOW_COLOR)
                    except (ValueError, TypeError):
                        pass
                elif i == status_col_offset:
                    cell.alignment = center
                    if val == "PASS":
                        cell.fill = mk_fill(PASS_COLOR); cell.font = mk_font(bold=True, color=PASS_FONT_COLOR)
                    else:
                        cell.fill = mk_fill(FAIL_COLOR); cell.font = mk_font(bold=True, color=FAIL_FONT_COLOR)
                else:
                    cell.font = mk_font(); cell.alignment = left
            ws2.row_dimensions[r].height = 28

        for i, w in enumerate(widths):
            ws2.column_dimensions[get_column_letter(start_col + i)].width = w

    # Security table is the size reference: 4 cols, widths [22, 26, 16, 12] = 76 total.
    # Eval / Ragas tables have 3 cols (Metric | Avg Score | Status); the Metric column
    # width is widened to 48 (=22+26) so each table's total width still sums to 76,
    # keeping all three summary tables the same overall size.
    build_summary_table(
        start_col=1, banner_text="EVALUATION SUMMARY",
        headers=["Metric", "Average Score", "Status"], rows=eval_summary,
        header_color=EVAL_HEADER_COLOR, widths=[48, 16, 12],
        score_col_offset=1, status_col_offset=2, score_threshold=0.7,
    )
    ws2.column_dimensions["D"].width = 4   # gap

    if ragas_summary:
        build_summary_table(
            start_col=5, banner_text="RAGAS EVALUATION SUMMARY",
            headers=["Metric", "Average Score", "Status"], rows=ragas_summary,
            header_color=RAGAS_HEADER_COLOR, widths=[48, 16, 12],
            score_col_offset=1, status_col_offset=2, score_threshold=0.7,
        )
    ws2.column_dimensions["H"].width = 4   # gap

    if sec_summary:
        build_summary_table(
            start_col=9, banner_text="SECURITY TEST SUMMARY",
            headers=["Vulnerability", "Vulnerability Type", "Average Score", "Status"], rows=sec_summary,
            header_color=SEC_HEADER_COLOR, widths=[22, 26, 16, 12],
            score_col_offset=2, status_col_offset=3, score_threshold=1.0,
        )

    wb.save(output_file)
    print(f"\n📊 Report saved as {output_file}")
